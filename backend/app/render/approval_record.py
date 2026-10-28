"""Выгружаемая запись согласования; не имитирует электронную подпись."""
import io
from datetime import datetime

from docx import Document
from docx.shared import Pt

from app.schemas.review import ApprovalDecision


def render_approval_record(session_id: str, decision: ApprovalDecision, current_revision: str) -> bytes:
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading('Лист согласования образовательной программы', 0)
    doc.add_heading(decision.program_title or 'Название не сохранено в записи решения', 1)
    rows = [
        ('Решение', 'Согласовано' if decision.decision == 'approved' else 'Возвращено на доработку'),
        ('Участник (имя введено пользователем)', decision.author),
        ('Дата решения (UTC)', datetime.fromisoformat(decision.at).strftime('%d.%m.%Y %H:%M:%S')),
        ('Идентификатор программы', session_id),
        ('Версия, к которой относится решение', decision.revision),
        ('Состояние версии на момент выгрузки', 'Текущая версия' if decision.revision == current_revision else 'Программа изменена после решения; требуется новое согласование'),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    doc.add_heading('Комментарий', 2)
    doc.add_paragraph(decision.comment or 'Без замечаний к решению.')
    if decision.review_summary:
        summary = decision.review_summary
        doc.add_heading('Состояние экспертизы на момент решения', 2)
        doc.add_paragraph(f"Согласовано материалов: {summary.get('reviewed', 0)} из {summary.get('total', 0)}.")
        doc.add_paragraph(f"Ожидали решения предложений: {summary.get('pending', 0)}. Материалов с замечаниями: {summary.get('needs_changes', 0)}.")
    doc.add_paragraph('Документ фиксирует согласование в системе. Имя участника введено самостоятельно. Электронная подпись отсутствует.')
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
