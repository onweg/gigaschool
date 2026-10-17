"""Выгрузка программы в .docx по шаблонам АЦР.

Структура разделов и формулировки титульного листа воспроизводят образцы
программ АЦР (data/acr/*.docx). Шаблон собирается кодом, а не правкой
заполненного образца: так в документ физически не может протечь чужой текст,
а состав разделов проверяется теми же тестами, что и остальной пайплайн.

Реквизиты приказа, ФИО преподавателей и подписи остаются пустыми —
их заполняет человек (требование кейса).
"""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.schemas.ir import (
    Program,
    ProgramType,
    StudyForm,
    VerificationStatus,
)

FONT = "Times New Roman"
SIZE = Pt(14)
SIZE_TABLE = Pt(10)

STUDY_FORM_RU = {
    StudyForm.FULL_TIME: "очная",
    StudyForm.PART_TIME: "очно-заочная",
    StudyForm.DISTANT: "заочная",
}

#: Падежные формы для шаблонных фраз («имеет техническую направленность»).
PROFILE_ACC = {
    "техническая": "техническую",
    "естественнонаучная": "естественнонаучную",
    "социально-гуманитарная": "социально-гуманитарную",
}
LEVEL_PREP = {"базовый": "базовом", "продвинутый": "продвинутом", "стартовый": "стартовом"}

PROGRAM_TITLE_RU = {
    ProgramType.DPP_PK: "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА\nПОВЫШЕНИЯ КВАЛИФИКАЦИИ",
    ProgramType.DOP_OBSH: "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ\nОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА",
}

UNVERIFIED_MARK = " [требует проверки методистом]"


def hours_str(value: float) -> str:
    """Часы в документе: целые — без дробной части, как в образцах АЦР."""
    if value is None:
        return "—"
    return f"{value:g}" if value else "0"


def _weeks_word(n: int) -> str:
    if n % 10 == 1 and n % 100 != 11:
        return "неделя"
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return "недели"
    return "недель"


class AcrDocxBuilder:
    def __init__(self, program: Program) -> None:
        self.p = program
        self.doc = Document()
        self._setup_styles()

    # ----------------------------- оформление --------------------------- #
    def _setup_styles(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = FONT
        style.font.size = SIZE
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        fmt = style.paragraph_format
        fmt.line_spacing = 1.5
        fmt.space_after = Pt(0)
        fmt.first_line_indent = Cm(1.25)
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for section in self.doc.sections:
            section.left_margin = Cm(3)
            section.right_margin = Cm(1.5)
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)

    def para(
        self,
        text: str = "",
        *,
        bold: bool = False,
        align: str = "justify",
        indent: bool = True,
        size: Pt | None = None,
        italic: bool = False,
    ):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.alignment = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        if not indent:
            paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT
        run.font.size = size or SIZE
        return paragraph

    def heading(self, text: str, *, level: int = 1) -> None:
        self.para("")
        self.para(text, bold=True, align="left", indent=False)

    def bullets(self, items: list[str], *, numbered: bool = False) -> None:
        for i, item in enumerate(items, start=1):
            prefix = f"{i}. " if numbered else "— "
            self.para(f"{prefix}{item}", align="justify")

    def table(
        self,
        rows: list[list[str]],
        *,
        widths: list[float] | None = None,
        header_rows: int = 1,
    ):
        if not rows:
            return None
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if r < header_rows else WD_ALIGN_PARAGRAPH.LEFT
                )
                run = paragraph.add_run(str(value))
                run.bold = r < header_rows
                run.font.name = FONT
                run.font.size = SIZE_TABLE
                if widths and c < len(widths):
                    cell.width = Cm(widths[c])
        return table

    def page_break(self) -> None:
        self.doc.add_page_break()

    # ------------------------------ титул -------------------------------- #
    def title_page(self) -> None:
        meta = self.p.meta
        for line in meta.organization.split("\n"):
            self.para(line, align="center", indent=False)
        self.para(
            "ОГРН: 1230600003457, ИНН/КПП: 0600010064/060001001",
            align="center",
            indent=False,
            size=Pt(12),
        )
        self.para(
            "Российская Федерация, 386001, Республика Ингушетия, "
            "городской округ Магас, город Магас, улица Н.С. Хрущева, дом 10",
            align="center",
            indent=False,
            size=Pt(12),
        )
        for _ in range(3):
            self.para("")
        self.para("УТВЕРЖДЕНО", align="right", indent=False)
        # Реквизиты приказа заполняет человек.
        self.para(
            "Приказом № ____-ОБР от «___» __________ "
            f"{meta.year} г.",
            align="right",
            indent=False,
        )
        for _ in range(4):
            self.para("")
        for line in PROGRAM_TITLE_RU[meta.program_type].split("\n"):
            self.para(line, bold=True, align="center", indent=False)
        self.para(f"«{meta.title}»", bold=True, align="center", indent=False)
        for _ in range(8):
            self.para("")
        self.para(f"{meta.city} {meta.year}", align="center", indent=False)
        self.page_break()

    # ------------------------- общие таблицы ----------------------------- #
    def study_plan_table(self) -> None:
        """2.1 Учебный план. Строка «Итого» считается из модулей, не из модели."""
        rows: list[list[str]] = [
            [
                "№ п/п",
                "Наименование разделов/модулей",
                "Общая трудоёмкость, час.",
                "ЛК",
                "ПР",
                "ЛБ",
                "СРС",
                "Форма промежуточной аттестации/контроля",
            ]
        ]
        for module in self.p.modules:
            rows.append(
                [
                    f"{module.index}.",
                    module.title,
                    hours_str(module.hours.total),
                    hours_str(module.hours.lectures),
                    hours_str(module.hours.practice),
                    hours_str(module.hours.labs),
                    hours_str(module.hours.self_study),
                    module.assessment_form or "—",
                ]
            )
        att = self.p.final_attestation
        rows.append(
            ["", "Итоговая аттестация", hours_str(att.hours), "—", "—", "—", "—", att.form]
        )
        total = self.p.modules_hours
        rows.append(
            [
                "",
                "Итого (час.)",
                hours_str(self.p.grand_total),
                hours_str(total.lectures),
                hours_str(total.practice),
                hours_str(total.labs),
                hours_str(total.self_study),
                "",
            ]
        )
        self.table(rows, widths=[1.2, 5.5, 1.8, 1.1, 1.1, 1.1, 1.2, 3.5])

    def topic_plan_table(self) -> None:
        """2.2 Учебно-тематический план: модули и темы с часами."""
        rows: list[list[str]] = [
            ["№", "Наименование разделов и тем", "Всего", "ЛК", "ПР", "ЛБ", "СРС", "Контроль"]
        ]
        for module in self.p.modules:
            rows.append(
                [
                    f"{module.index}.",
                    module.title,
                    hours_str(module.hours.total),
                    hours_str(module.hours.lectures),
                    hours_str(module.hours.practice),
                    hours_str(module.hours.labs),
                    hours_str(module.hours.self_study),
                    module.assessment_form or "",
                ]
            )
            for topic in module.topics:
                rows.append(
                    [
                        f"{topic.code}.",
                        topic.title,
                        hours_str(topic.hours.total),
                        hours_str(topic.hours.lectures),
                        hours_str(topic.hours.practice),
                        hours_str(topic.hours.labs),
                        hours_str(topic.hours.self_study),
                        topic.assessment_form or "",
                    ]
                )
        att = self.p.final_attestation
        rows.append(["", "Итоговая аттестация", hours_str(att.hours), "—", "—", "—", "—", att.form])
        total = self.p.modules_hours
        rows.append(
            [
                "",
                "Итого (час.)",
                hours_str(self.p.grand_total),
                hours_str(total.lectures),
                hours_str(total.practice),
                hours_str(total.labs),
                hours_str(total.self_study),
                "",
            ]
        )
        self.table(rows, widths=[1.2, 6.0, 1.3, 1.0, 1.0, 1.0, 1.1, 3.0])

    def outcomes_list(self) -> None:
        for kind, label in (
            ("знание", "Знать:"),
            ("умение", "Уметь:"),
            ("навык", "Владеть навыками:"),
        ):
            items = [o for o in self.p.outcomes if o.kind.value == kind]
            if not items:
                continue
            self.para(label, bold=True, align="left", indent=False)
            for o in items:
                self.para(f"— {o.text} ({o.id})")

    def working_program_table(self) -> None:
        """2.4 Рабочая программа: содержание и формы работы по темам."""
        rows: list[list[str]] = [
            ["Наименование компонентов программы", "Содержание и формы организации деятельности", "Всего, час."]
        ]
        for module in self.p.modules:
            rows.append([f"Раздел {module.index}. {module.title}", "", hours_str(module.hours.total)])
            for topic in module.topics:
                body = topic.content or "Содержание темы формируется на этапе наполнения."
                if topic.activities:
                    body = f"{body}\n{topic.activities}"
                rows.append([f"Тема {topic.code}. {topic.title}", body, hours_str(topic.hours.total)])
        rows.append(
            [
                "Итоговая аттестация",
                self.p.final_attestation.description or self.p.final_attestation.form,
                hours_str(self.p.final_attestation.hours),
            ]
        )
        rows.append(["Итого", "", hours_str(self.p.grand_total)])
        self.table(rows, widths=[5.0, 9.0, 1.8])

    def staffing_table(self) -> None:
        """Кадровое обеспечение: ФИО заполняет человек."""
        rows: list[list[str]] = [["№", "Модуль в соответствии с УП", "ФИО преподавателя", "Уч. степень/звание", "Должность"]]
        for module in self.p.modules:
            rows.append([str(module.index), f"Модуль {module.index}. {module.title}", "", "", ""])
        rows.append([str(len(self.p.modules) + 1), "Итоговая аттестация", "", "", ""])
        self.table(rows, widths=[1.0, 7.0, 3.5, 2.5, 2.5])

    def sources_block(self, prefix: str) -> None:
        groups = [
            ("main", f"{prefix}.1. Основная литература"),
            ("extra", f"{prefix}.2. Дополнительная литература"),
            ("internet", f"{prefix}.3. Интернет-ресурсы"),
            ("software", f"{prefix}.4. Программное обеспечение"),
        ]
        for kind, title in groups:
            items = [s for s in self.p.sources if s.kind == kind]
            if not items:
                continue
            self.heading(title)
            for i, source in enumerate(items, start=1):
                mark = (
                    UNVERIFIED_MARK
                    if source.verification is not VerificationStatus.VERIFIED
                    else ""
                )
                url = f" — URL: {source.url}" if source.url else ""
                self.para(f"{i}. {source.citation}{url}{mark}")

    def schedule_block(self) -> None:
        s = self.p.schedule
        if s.weeks:
            self.para(f"Срок обучения: {s.weeks} {_weeks_word(s.weeks)}.")
        elif s.days:
            self.para(f"Срок обучения: {s.days} дней.")
        if s.sessions and s.hours_per_session:
            self.para(
                f"Режим занятий: {s.sessions} занятий по {hours_str(s.hours_per_session)} "
                "академических часа."
            )
        if s.max_hours_per_day:
            self.para(f"Режим занятий: не более {hours_str(s.max_hours_per_day)} часов в день.")
        self.para(
            "Календарный учебный график формируется по мере комплектования учебных групп."
        )

    def assessment_block(self) -> None:
        att = self.p.final_attestation
        self.para(f"Форма итоговой аттестации: {att.form}.")
        if att.description:
            self.para(att.description)
        if att.criteria:
            self.para("Критерии оценки:", bold=True, align="left", indent=False)
            self.bullets(att.criteria)
        current = [m.assessment_form for m in self.p.modules if m.assessment_form]
        if current:
            self.para("Формы текущего контроля по модулям:", bold=True, align="left", indent=False)
            self.bullets(current)
        if self.p.test_questions:
            self.para(
                f"Банк тестовых вопросов: {len(self.p.test_questions)} вопросов с ключами "
                "(приложение «Тестовые задания»)."
            )
        if self.p.seminar_tasks:
            self.para(
                f"Задания для семинаров: {len(self.p.seminar_tasks)} заданий с критериями "
                "оценки (приложение «Задания для семинаров»)."
            )

    def appendix_tests(self) -> None:
        if not self.p.test_questions:
            return
        self.page_break()
        self.para("Приложение 1", align="right", indent=False)
        self.para("Тестовые задания с ключами", bold=True, align="center", indent=False)
        for i, q in enumerate(self.p.test_questions, start=1):
            self.para(f"{i}. {q.stem}", align="left")
            for j, option in enumerate(q.options):
                self.para(f"   {chr(1072 + j)}) {option}", align="left")
            keys = ", ".join(chr(1072 + k) for k in q.correct)
            trace = f"; проверяет: {', '.join(q.outcome_ids)}" if q.outcome_ids else ""
            self.para(f"   Ключ: {keys}{trace}", italic=True, align="left")

    def appendix_seminars(self) -> None:
        if not self.p.seminar_tasks:
            return
        self.page_break()
        self.para("Приложение 2", align="right", indent=False)
        self.para("Задания для семинаров", bold=True, align="center", indent=False)
        for i, task in enumerate(self.p.seminar_tasks, start=1):
            self.para(f"{i}. {task.title}", bold=True, align="left")
            self.para(task.statement)
            if task.deliverable:
                self.para(f"Результат: {task.deliverable}")
            if task.criteria:
                self.para("Критерии оценки:", align="left", indent=False)
                self.bullets(task.criteria)
            if task.outcome_ids:
                self.para(f"Проверяет: {', '.join(task.outcome_ids)}", italic=True)

    def traceability_appendix(self) -> None:
        """Карта компетенций: РО → темы → оценочные средства.

        Отдельное приложение, потому что проверяющему нужно за один взгляд
        убедиться, что каждая компетенция действительно проверяется.
        """
        if not self.p.outcomes:
            return
        self.page_break()
        self.para("Приложение 3", align="right", indent=False)
        self.para(
            "Карта компетенций и результатов обучения", bold=True, align="center", indent=False
        )
        rows: list[list[str]] = [
            ["Результат обучения", "Вид", "Темы", "Оценочные средства"]
        ]
        for outcome in self.p.outcomes:
            topics = [t.code for t in self.p.all_topics() if outcome.id in t.outcome_ids]
            tools: list[str] = []
            if any(outcome.id in q.outcome_ids for q in self.p.test_questions):
                tools.append("тест")
            if any(outcome.id in s.outcome_ids for s in self.p.seminar_tasks):
                tools.append("семинар")
            if self.p.final_attestation.form:
                tools.append(self.p.final_attestation.form)
            rows.append(
                [
                    f"{outcome.id}. {outcome.text}",
                    outcome.kind.value,
                    ", ".join(topics) or "—",
                    ", ".join(tools) or "—",
                ]
            )
        self.table(rows, widths=[7.5, 2.0, 2.5, 4.0])

    # --------------------------- шаблон ДПП ПК --------------------------- #
    def build_dpp_pk(self) -> None:
        p = self.p
        self.title_page()
        self.para("1. ОБЩАЯ ХАРАКТЕРИСТИКА ПРОГРАММЫ", bold=True, align="center", indent=False)
        self.heading("1.1. Аннотация программы")
        self.para(p.annotation or p.goal)
        self.heading("1.2. Цель реализации программы")
        self.para(p.goal)
        if p.objectives:
            self.para("Задачи программы:", align="left", indent=False)
            self.bullets(p.objectives, numbered=True)
        self.heading("1.3. Планируемые результаты обучения (РО)")
        self.outcomes_list()
        if p.competencies:
            rows: list[list[str]] = [
                ["Виды деятельности", "Профессиональные компетенции", "Результаты обучения"]
            ]
            for c in p.competencies:
                linked = [o.text for o in p.outcomes if o.id in c.outcome_ids]
                rows.append([c.activity or "—", c.title, "; ".join(linked) or "—"])
            self.table(rows, widths=[5.0, 5.5, 5.5])
        self.heading("1.4. Программа разработана на основе:")
        for i, act in enumerate(p.legal_acts, start=1):
            mark = UNVERIFIED_MARK if act.verification is not VerificationStatus.VERIFIED else ""
            self.para(f"{i}. {act.requisites} «{act.title}».{mark}")
        for std in p.prof_standards:
            mark = (
                UNVERIFIED_MARK
                if std.verification is not VerificationStatus.VERIFIED
                else ""
            )
            approved = f" (утв. {std.approved_by})" if std.approved_by else ""
            self.para(
                f"Профессионального стандарта {std.code} «{std.title}»{approved}.{mark}"
            )
            if std.note:
                self.para(f"Примечание: {std.note}", italic=True)
        self.heading(
            "1.5. Перечень профессиональных компетенций, качественное изменение "
            "которых осуществляется в результате обучения:"
        )
        for c in p.competencies:
            lf = f"{c.labor_function}. " if c.labor_function else ""
            self.para(f"— {lf}{c.title}")
        self.heading("1.6. Требования к слушателям (категории слушателей)")
        self.para(p.meta.audience or "—")
        if p.entry_requirements:
            self.para(p.entry_requirements)
        self.heading("1.7. Форма обучения")
        form = STUDY_FORM_RU.get(p.meta.study_form, "очно-заочная")
        dot = " с применением электронного обучения и дистанционных образовательных технологий" if p.meta.uses_dot else ""
        self.para(f"{form.capitalize()}{dot}.")
        self.heading("1.8. Объём программы")
        self.para(f"{hours_str(p.meta.total_hours)} академических часов.")
        self.heading("1.9. Выдаваемый документ")
        self.para(
            f"Лицам, успешно освоившим программу и прошедшим итоговую аттестацию, "
            f"выдаётся {p.meta.issued_document}."
        )
        self.page_break()
        self.para("2. СОДЕРЖАНИЕ ПРОГРАММЫ", bold=True, align="center", indent=False)
        self.heading("2.1. Учебный план")
        self.study_plan_table()
        self.heading("2.2. Учебно-тематический план")
        self.topic_plan_table()
        self.heading("2.3. Календарный учебный график")
        self.schedule_block()
        self.heading("2.4. Рабочая программа учебных дисциплин")
        self.working_program_table()
        self.heading("2.5. Оценка качества освоения программы")
        self.assessment_block()
        self.page_break()
        self.para(
            "3. ОРГАНИЗАЦИОННО-ПЕДАГОГИЧЕСКИЕ УСЛОВИЯ РЕАЛИЗАЦИИ ПРОГРАММЫ",
            bold=True,
            align="center",
            indent=False,
        )
        self.heading("3.1. Материально-техническое обеспечение")
        self.bullets(p.material_conditions or ["учебная аудитория с проектором", "компьютеры с доступом в интернет"], numbered=True)
        self.heading("3.2. Учебно-методическое и информационное обеспечение")
        self.heading("3.2.1. Нормативно-правовые документы")
        for i, act in enumerate(p.legal_acts, start=1):
            self.para(f"{i}. {act.requisites} «{act.title}».")
        self.sources_block("3.2")
        self.heading("3.3. Кадровое обеспечение программы")
        self.staffing_table()
        self.appendix_tests()
        self.appendix_seminars()
        self.traceability_appendix()

    # ------------------- шаблон общеразвивающей программы ---------------- #
    def build_dop_obsh(self) -> None:
        p = self.p
        self.title_page()
        self.para("Пояснительная записка", bold=True, align="center", indent=False)
        profile = PROFILE_ACC.get(p.meta.profile, p.meta.profile)
        level = p.meta.level or "базовый"
        self.para(
            f"Дополнительная общеразвивающая программа «{p.meta.title}» имеет "
            f"{profile} направленность и реализуется на "
            f"{LEVEL_PREP.get(level, level)} уровне."
        )
        if p.annotation:
            self.para(p.annotation)
        self.para(f"Уровень освоения программы: {p.meta.level or 'базовый'}.")
        self.para(f"Направленность (профиль) программы — {p.meta.profile}.")
        form = STUDY_FORM_RU.get(p.meta.study_form, "очная")
        self.para(f"Форма обучения: {form}.")
        self.heading("1.1. Актуальность программы")
        self.para(p.relevance or p.annotation or p.goal)
        self.heading("1.2. Категории обучающихся")
        self.para(p.meta.audience or "—")
        if p.entry_requirements:
            self.para(p.entry_requirements)
        self.heading("1.3. Цель и задачи программы")
        self.para(f"Цель программы: {p.goal}")
        if p.objectives:
            self.para("Задачи программы:", align="left", indent=False)
            self.bullets(p.objectives, numbered=True)
        self.heading("1.4. Сроки реализации")
        self.para(f"Объём программы: {hours_str(p.meta.total_hours)} академических часов.")
        self.schedule_block()

        self.heading("1.5. Планируемые результаты обучения")
        self.outcomes_list()
        self.heading("1.6. Календарный учебный график")
        self.para(
            "Календарный учебный график формируется по мере комплектования учебных "
            "групп и утверждается отдельным документом. Распределение часов по "
            "занятиям приведено в учебно-тематическом плане (раздел 2.2)."
        )
        self.page_break()
        self.para("2. СОДЕРЖАНИЕ ПРОГРАММЫ", bold=True, align="center", indent=False)
        self.heading("2.1. Учебный план")
        self.study_plan_table()
        self.heading("2.2. Учебно-тематический план")
        self.topic_plan_table()
        self.heading("2.3. Рабочая программа")
        self.working_program_table()
        self.heading("2.4. Оценочные материалы")
        self.assessment_block()
        self.heading("2.5. Методические материалы")
        self.para(
            "Обучение строится на разборе практических задач: теоретический блок "
            "сопровождается практикой, по итогам модулей проводится текущий контроль."
        )
        self.page_break()
        self.para(
            "3. ОРГАНИЗАЦИОННО-ПЕДАГОГИЧЕСКИЕ УСЛОВИЯ РЕАЛИЗАЦИИ ПРОГРАММЫ",
            bold=True,
            align="center",
            indent=False,
        )
        self.heading("3.1. Материально-технические условия реализации программы")
        self.bullets(p.material_conditions or ["учебная аудитория с проектором", "компьютеры с доступом в интернет"], numbered=True)
        self.heading("3.2. Учебно-методическое и информационное обеспечение программы")
        for i, act in enumerate(p.legal_acts, start=1):
            self.para(f"{i}. {act.requisites} «{act.title}».")
        self.sources_block("3.3")
        self.heading("3.4. Кадровое обеспечение программы")
        self.staffing_table()
        self.appendix_tests()
        self.appendix_seminars()
        self.traceability_appendix()

    # ------------------------------ сборка ------------------------------- #
    def build(self) -> Document:
        if self.p.meta.program_type is ProgramType.DPP_PK:
            self.build_dpp_pk()
        else:
            self.build_dop_obsh()
        return self.doc


def render_docx(program: Program) -> bytes:
    buffer = BytesIO()
    AcrDocxBuilder(program).build().save(buffer)
    return buffer.getvalue()
