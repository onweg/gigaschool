"""Выгрузка слайдов лекций в .pptx.

Слайды собираются из LecturePlan: заголовок + тезисы, заметки докладчика
уходят в notes slide. Оформление сдержанное — на защите важна согласованность
с учебным планом, а не дизайн.
"""
from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

from app.schemas.ir import LecturePlan, Program

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)


def _title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = ACCENT
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def _content_slide(prs: Presentation, title: str, bullets: list[str], notes: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(30)
            run.font.color.rgb = ACCENT
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, bullet in enumerate(bullets or ["—"]):
        paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        for run in paragraph.runs:
            run.font.size = Pt(20)
            run.font.color.rgb = TEXT
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def render_pptx(program: Program, plans: list[LecturePlan] | None = None) -> bytes:
    plans = plans if plans is not None else program.lecture_plans
    prs = Presentation()
    prs.slide_width = Cm(33.87)
    prs.slide_height = Cm(19.05)

    _title_slide(
        prs,
        program.meta.title,
        f"{program.meta.total_hours:g} ак. ч · {program.meta.audience or 'слушатели'}",
    )
    for plan in plans:
        _title_slide(prs, plan.title, "; ".join(plan.objectives[:2]))
        for slide in plan.slides:
            _content_slide(prs, slide.title, slide.bullets, slide.notes)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
