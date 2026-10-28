"""Экспертиза добровольна; принятые правки версионны и доходят до выгрузки."""
import io
import json
from copy import deepcopy

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.llm.base import Completion, LLMProvider, Usage
from app.llm.mock import MockProvider
from app.pipeline.expert_review import analyze, material_items, suggest
from app.store import Stage, SessionStore, store
from conftest import READY

client = TestClient(app)


@pytest.fixture
def ready():
    r = client.post('/api/generate', json={
        'text': 'Повышение квалификации по Python на 72 часа для сотрудников банков, очно-заочно, защита проекта.',
        'answers': READY,
    })
    assert r.status_code == 200, r.text
    sid = r.json()['session_id']
    return sid, f'/api/sessions/{sid}/review'


def view(base):
    r = client.get(base)
    assert r.status_code == 200, r.text
    return r.json()


def proposal(base, item, suffix=' Эксперт уточнил материал.'):
    after = deepcopy(item['value'])
    field = 'stem' if item['kind'] == 'test' else 'content'
    after[field] += suffix
    r = client.post(base + '/proposals', json={'item_id': item['id'], 'revision': item['revision'],
        'after': after, 'author': 'Анна, эксперт', 'comment': 'Уточнить формулировку'})
    assert r.status_code == 200, r.text
    return r.json()['proposals'][-1]


def test_solo_export_and_approval_do_not_require_expert(ready):
    sid, base = ready
    d = view(base)
    assert d['summary']['reviewed'] == 0
    assert client.get(f'/api/sessions/{sid}/export/docx').status_code == 200
    r = client.post(base + '/approval', json={'revision': d['revision'], 'author': 'Руководитель', 'decision': 'approved'})
    assert r.status_code == 200, r.text
    assert r.json()['approval']['decision'] == 'approved'
    assert not r.json()['approval']['stale']
    assert store.get(sid).stage == Stage.READY


def test_proposal_is_not_applied_until_methodist_accepts_and_export_reflects_it(ready):
    sid, base = ready
    item = next(i for i in view(base)['items'] if i['kind'] == 'test')
    original = client.get(f'/api/sessions/{sid}').json()['program']
    p = proposal(base, item)
    assert client.get(f'/api/sessions/{sid}').json()['program'] == original
    assert client.get(f'/api/sessions/{sid}/export/docx').status_code == 200
    response = client.post(base + f"/proposals/{p['id']}/resolve", json={'action': 'accept'})
    assert response.status_code == 200, response.text
    next_program = response.json()['session']['program']
    assert next_program['modules'] == original['modules']
    assert next_program['hours_check']['ok']
    assert next_program['test_questions'][0]['stem'] == p['after']['stem']
    assert all(q1 == q2 for q1, q2 in zip(original['test_questions'][1:], next_program['test_questions'][1:]))
    exported = client.get(f'/api/sessions/{sid}/export/csv')
    assert 'Эксперт уточнил материал.' in exported.text
    doc = Document(io.BytesIO(client.get(f'/api/sessions/{sid}/export/docx').content))
    text = '\n'.join(p.text for p in doc.paragraphs)
    text += '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'Эксперт уточнил материал.' in text
    assert client.post(base + f"/proposals/{p['id']}/resolve", json={'action':'accept'}).status_code == 409


def test_topic_change_invalidates_related_test_review_and_approval(ready):
    sid, base = ready
    d = view(base)
    topic = d['items'][0]
    test = next(i for i in d['items'] if i['kind'] == 'test' and i['topic_id'] == topic['topic_id'])
    client.post(base + '/decisions', json={'item_id': test['id'], 'revision': test['revision'], 'status': 'reviewed', 'author': 'Эксперт'})
    client.post(base + '/approval', json={'revision': d['revision'], 'author': 'Руководитель', 'decision': 'approved'})
    p = proposal(base, topic)
    client.post(base + f"/proposals/{p['id']}/resolve", json={'action':'accept'})
    d = view(base)
    assert next(i for i in d['items'] if i['id'] == test['id'])['status'] == 'stale'
    assert d['approval']['stale']
    assert client.post(base + '/decisions', json={'item_id': test['id'], 'revision': test['revision'], 'status':'reviewed', 'author':'Эксперт'}).status_code == 409


def test_conflicting_proposals_cannot_overwrite_and_can_be_rejected(ready):
    sid, base = ready
    item = view(base)['items'][0]
    a = proposal(base, item, ' Вариант А.')
    b = proposal(base, item, ' Вариант Б.')
    client.post(base + f"/proposals/{a['id']}/resolve", json={'action':'accept'})
    r = client.post(base + f"/proposals/{b['id']}/resolve", json={'action':'accept'})
    assert r.status_code == 409
    assert view(base)['proposals'][-1]['stale']
    r = client.post(base + f"/proposals/{b['id']}/resolve", json={'action':'reject'})
    assert r.status_code == 200
    assert r.json()['session']['program']['modules'][0]['topics'][0]['content'] == a['after']['content']


@pytest.mark.parametrize('bad', [ {'correct':[99]}, {'correct':[]}, {'correct':[0,0]}, {'correct':[True]}, {'options':['same','same']}, {'hours':72}, {'id':'other'} ])
def test_invalid_edits_are_rejected_without_changing_program(ready, bad):
    sid, base = ready
    d = view(base)
    item = next(i for i in d['items'] if i['kind'] == 'test')
    r = client.post(base + '/proposals', json={'item_id':item['id'], 'revision':item['revision'], 'author':'Эксперт', 'after':{**item['value'], **bad}})
    assert r.status_code == 422, r.text
    assert view(base)['revision'] == d['revision']
    assert not view(base)['proposals']


def test_review_progress_persists_and_opening_does_not_mark_reviewed(ready, tmp_path):
    sid, base = ready
    d = view(base)
    assert all(i['status'] == 'todo' for i in d['items'])
    item = d['items'][0]
    client.post(base + '/decisions', json={'item_id':item['id'],'revision':item['revision'], 'author':'Анна','status':'deferred'})
    persisted = SessionStore(str(tmp_path))
    persisted.save(store.get(sid))
    loaded = SessionStore(str(tmp_path)).get(sid)
    assert loaded.review.decisions[item['id']].status == 'deferred'
    assert view(base)['summary']['reviewed'] == 0


def test_review_blocked_during_generation_and_comments_required(ready):
    sid, base = ready
    d = view(base)
    item = d['items'][0]
    payload = {'item_id':item['id'], 'revision':item['revision'], 'author':'Анна', 'status':'needs_changes'}
    assert client.post(base + '/decisions', json=payload).status_code == 422
    payload['status'] = 'reviewed'
    store.get(sid).stage = Stage.CONTENT
    assert client.post(base + '/decisions', json=payload).status_code == 409
    # Даже незавершённая экспертиза не становится условием экспорта.
    assert client.get(f'/api/sessions/{sid}/export/docx').status_code == 200


def test_scan_does_not_claim_review_and_mock_edit_is_honest(ready):
    sid, base = ready
    item = view(base)['items'][0]
    r = client.post(base + '/scan', json={'topic_id':item['topic_id']})
    assert r.status_code == 200, r.text
    assert r.json()['scans'][item['topic_id']]['offline']
    assert r.json()['summary']['reviewed'] == 0
    assert not any(i['findings'] for i in r.json()['items'])
    r = client.post(base + '/suggest', json={'item_id':item['id'], 'revision':item['revision'], 'instruction':'Уточнить пример'})
    assert r.status_code == 200, r.text
    assert r.json()['offline']
    assert r.json()['replacement'] == item['value']
    assert not view(base)['proposals']


@pytest.mark.anyio
async def test_agent_findings_are_grounded_and_suggestion_does_not_mutate(ready):
    sid, _ = ready
    program = store.get(sid).program
    item = material_items(program)[0]
    quote = item['value']['content'][:40]
    class Spy(LLMProvider):
        name = model = 'spy'
        async def complete(self, *, system, user, **kwargs):
            assert program.meta.audience in user
            assert item['value']['content'] in user or json.dumps(item['value']['content'], ensure_ascii=False)[1:-1] in user
            if 'SKILL:expert_review' in system:
                payload = {'findings':[
                    {'item_id':item['id'], 'title':'Уточнить', 'reason':'Требуется контекст', 'quote':quote, 'suggestion':'Добавить пример'},
                    {'item_id':'missing', 'title':'Нет', 'reason':'Нет', 'quote':quote},
                    {'item_id':item['id'], 'title':'Выдумка', 'reason':'Нет', 'quote':'ЭТОГО НЕТ В МАТЕРИАЛЕ'},
                ]}
            else:
                assert 'Добавить пример' in user
                payload = {'replacement': {**item['value'], 'content':item['value']['content']+' Пример.'}}
            return Completion(text=json.dumps(payload, ensure_ascii=False), usage=Usage())
    before = program.model_dump()
    findings, _ = await analyze(Spy(), program, [item])
    assert len(findings) == 1
    assert findings[0].quote == quote
    replacement, _ = await suggest(Spy(), program, item, 'Добавить пример')
    assert replacement['content'].endswith('Пример.')
    assert program.model_dump() == before


def test_grounded_finding_can_be_dismissed_without_marking_material_reviewed(ready, monkeypatch):
    from app.schemas.review import ReviewFinding
    sid, base = ready
    item = view(base)['items'][0]
    async def fake_analyze(provider, program, items):
        return [ReviewFinding(item_id=item['id'], revision=item['revision'], title='Уточнить пример',
            reason='Недостаточно контекста', quote=item['value']['content'][:30], suggestion='Добавить пример')], Usage()
    monkeypatch.setattr('app.routes.review.analyze', fake_analyze)
    r = client.post(base + '/scan', json={'topic_id': item['topic_id']})
    assert r.status_code == 200
    finding = r.json()['items'][0]['findings'][0]
    r = client.post(base + f"/findings/{finding['id']}/dismiss", json={'revision':item['revision']})
    assert r.status_code == 200
    assert r.json()['items'][0]['findings'][0]['dismissed']
    assert r.json()['items'][0]['status'] == 'todo'


def test_analysis_does_not_attach_findings_to_changed_material(ready, monkeypatch):
    sid, base = ready
    item = view(base)['items'][0]
    async def simultaneous_edit(provider, program, items):
        store.get(sid).program.all_topics()[0].content += ' Изменено другим участником.'
        return [], Usage()
    monkeypatch.setattr('app.routes.review.analyze', simultaneous_edit)
    r = client.post(base + '/scan', json={'topic_id':item['topic_id']})
    assert r.status_code == 409
    assert not view(base)['scans']


def test_batch_approves_only_selected_without_changing_content(ready):
    sid, base = ready
    before = view(base)
    selected = before['items'][:3]
    r = client.post(base + '/decisions/batch', json={'author':'Эксперт', 'items':[
        {'item_id':i['id'], 'revision':i['revision']} for i in selected]})
    assert r.status_code == 200, r.text
    assert r.json()['summary']['reviewed'] == 3
    assert all(i['decision']['method'] == 'batch' for i in r.json()['items'][:3])
    assert all(i['status'] == 'todo' for i in r.json()['items'][3:])
    assert r.json()['revision'] == before['revision']
    assert client.get(f'/api/sessions/{sid}/export/docx').status_code == 200


def test_batch_is_atomic_if_one_revision_changed(ready):
    sid, base = ready
    items = view(base)['items'][:3]
    payload = [{'item_id':i['id'], 'revision':i['revision']} for i in items]
    payload[-1]['revision'] = 'old'
    r = client.post(base + '/decisions/batch', json={'author':'Эксперт','items':payload})
    assert r.status_code == 409
    assert view(base)['summary']['reviewed'] == 0
    assert not store.get(sid).review.decision_history


@pytest.mark.parametrize('items,author', [([], 'Эксперт'), ([{'item_id':'missing','revision':'old'}], ' '),
    ([{'item_id':'same','revision':'old'}, {'item_id':'same','revision':'old'}], 'Эксперт')])
def test_batch_rejects_empty_duplicate_or_unnamed_selection(ready, items, author):
    _, base = ready
    assert client.post(base + '/decisions/batch', json={'items':items, 'author':author}).status_code == 422


def test_batch_keeps_individual_decisions_proposals_and_comment_history(ready):
    _, base = ready
    items = view(base)['items']
    first = items[0]
    client.post(base + '/decisions', json={'item_id':first['id'], 'revision':first['revision'], 'author':'Первый эксперт', 'status':'reviewed'})
    p = proposal(base, items[1])
    r = client.post(base + '/decisions/batch', json={'author':'Второй эксперт','items':[
        {'item_id':i['id'],'revision':i['revision']} for i in items]})
    assert r.status_code == 200
    d = r.json()
    assert d['summary']['reviewed'] == len(items)
    assert d['items'][0]['decision']['method'] == 'individual'
    assert d['items'][0]['decision']['author'] == 'Первый эксперт'
    assert d['proposals'][0]['id'] == p['id'] and d['proposals'][0]['status'] == 'pending'
    assert any(h['comment'] == 'Уточнить формулировку' for h in d['decision_history'])


def test_approval_record_preserves_decision_snapshot_and_marks_later_changes(ready):
    sid, base = ready
    assert client.get(base + '/approval-record.docx').status_code == 409
    d = view(base)
    title = store.get(sid).program.meta.title
    r = client.post(base + '/approval', json={'revision':d['revision'],'author':'Руководитель','decision':'approved','comment':'К запуску готово'})
    assert r.status_code == 200
    store.get(sid).program.meta.title = 'Изменённое название после согласования'
    r = client.get(base + '/approval-record.docx')
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for value in [title, d['revision'], 'Руководитель', 'К запуску готово', 'Программа изменена после решения', 'Электронная подпись отсутствует']:
        assert value in text
    assert 'Изменённое название после согласования' not in text
    assert 'Согласовано материалов: 0' in text


def test_question_to_agent_never_changes_program_or_approves_material(ready):
    _, base = ready
    d = view(base)
    item = d['items'][1]
    r = client.post(base + '/ask', json={'item_id':item['id'], 'revision':item['revision'], 'instruction':'Почему верен первый вариант?'})
    assert r.status_code == 200, r.text
    assert 'Офлайн-режим' in r.json()['answer']
    after = view(base)
    assert after['revision'] == d['revision']
    assert after['summary']['reviewed'] == 0
    assert after['proposals'] == []
