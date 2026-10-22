"""Замечания автопроверки доходят до выгрузок.

Проверяющие отметили: «экспорт не включает явных пояснений о серьёзности
замечаний автопроверки и статусах». Методист открывает .docx отдельно от
интерфейса и должен видеть, что именно проверить перед утверждением.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from conftest import READY

from app.core.validate import validate_program
from app.llm.mock import MockProvider
from app.main import app
from app.pipeline.skeleton import build_skeleton
from app.render.docx_acr import render_docx
from app.schemas.ir import Program, ProgramMeta, ProgramType
from app.schemas.request import ProgramRequest

client = TestClient(app)


def _text(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.anyio
async def test_appendix_lists_findings_with_what_to_do() -> None:
    request = ProgramRequest(
        title="Анализ данных",
        program_type=ProgramType.DPP_PK,
        total_hours=72,
        audience="сотрудники банков",
        entry_level="базовый",
    )
    program, _ = await build_skeleton(MockProvider(), request)
    report = validate_program(program, stage="skeleton")

    text = _text(render_docx(program, report))

    assert "ПРИЛОЖЕНИЕ. ОТЧЁТ АВТОПРОВЕРКИ" in text
    assert "Что сделать" in text
    for finding in report.findings:
        assert finding.message[:40] in text, f"нет находки {finding.code}"
        assert finding.hint, f"у {finding.code} нет подсказки методисту"
        assert finding.hint[:30] in text, f"нет подсказки для {finding.code}"


def test_every_check_explains_what_to_do() -> None:
    """Код находки методисту ничего не говорит — нужна подсказка словами."""
    empty = Program(
        meta=ProgramMeta(title="Пустая", program_type=ProgramType.DPP_PK, total_hours=72)
    )
    report = validate_program(empty)

    assert report.findings, "пустая программа обязана дать находки"
    without = [f.code for f in report.findings if not f.hint]
    assert not without, f"без подсказки: {without}"


@pytest.mark.anyio
async def test_clean_program_says_so_explicitly() -> None:
    """Пустое приложение хуже отсутствующего: методист не поймёт, проверяли ли."""
    request = ProgramRequest(
        title="Чистая",
        program_type=ProgramType.DPP_PK,
        total_hours=72,
        audience="слушатели",
        entry_level="базовый",
    )
    program, _ = await build_skeleton(MockProvider(), request)

    class Clean:
        findings: list = []

    text = _text(render_docx(program, Clean()))

    assert "замечаний не выявила" in text
    assert "подписи оставлены" in text


def test_json_export_carries_a_summary() -> None:
    body = client.post(
        "/api/generate",
        json={"text": "программа повышения квалификации на 72 часа по анализу данных "
                      "для банков, очно-заочно, защита проекта", "answers": READY},
    ).json()
    assert body["status"] == "ready", body

    bundle = client.get(body["exports"]["json"]).json()
    summary = bundle["validation_summary"]

    assert set(summary) >= {"ok", "errors", "warnings", "must_fix", "check_yourself", "legend"}
    assert isinstance(summary["errors"], int)
    # Легенда объясняет, что значит каждая степень — без неё числа бесполезны.
    assert "утверждена" in summary["legend"]["errors"]
    for item in summary["check_yourself"]:
        assert item["hint"], f"в сводке нет подсказки для {item['code']}"


def test_docx_export_includes_the_appendix() -> None:
    body = client.post(
        "/api/generate",
        json={"text": "общеразвивающая программа на 16 часов для школьников, очно",
              "answers": READY},
    ).json()
    assert body["status"] == "ready", body

    docx = client.get(body["exports"]["docx"])
    assert docx.status_code == 200
    assert "ОТЧЁТ АВТОПРОВЕРКИ" in _text(docx.content)
