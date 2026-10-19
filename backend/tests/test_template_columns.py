"""Колонки часов по шаблону АЦР.

В образцах (data/acr/) они разные:
  повышение квалификации — ЛК · ПР · ЛБ · СРС;
  общеразвивающая — всего аудиторных · лекции · практика · самостоятельная,
  лабораторных нет вовсе.

Команда заметила, что в общеразвивающей программе стояла колонка ЛБ, которой
нет в образце. Часы лабораторных уходят в практику, а не в самостоятельную
работу: лабораторная — занятие в классе, и иначе развалится расписание
«8 занятий по 2 часа» из чек-листа запроса 1.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.llm.mock import MockProvider
from app.pipeline.intake import apply_answers, run_intake
from app.pipeline.skeleton import assemble_program, build_skeleton
from app.render.docx_acr import render_docx
from app.schemas.ir import ProgramType, StudyForm
from app.schemas.request import ProgramRequest

REQUEST_1 = (
    "Нужна дополнительная общеразвивающая программа «Первые шаги в Python» для "
    "школьников 8 и 9 классов без опыта программирования. Объем 16 часов, занятия "
    "очно по 2 часа два раза в неделю. Практики должно быть не меньше половины всех часов."
)

#: Ответ модели, где лабораторным явно отдана доля — как бывает на живой модели.
WITH_LABS = {
    "goal": "Цель",
    "outcomes": [{"id": "RO-1", "kind": "умение", "text": "уметь писать программы"}],
    "modules": [
        {
            "title": "Основы",
            "weight": 1,
            "mix": {"lectures": 1, "practice": 2, "labs": 2, "self_study": 0},
            "outcome_ids": ["RO-1"],
            "topics": [{"title": "Переменные"}, {"title": "Циклы"}],
        }
    ],
}


def _request(ptype: ProgramType) -> ProgramRequest:
    return ProgramRequest(
        title="Т",
        program_type=ptype,
        total_hours=36,
        audience="слушатели",
        entry_level="базовый",
        study_form=StudyForm.FULL_TIME,
        attestation_form="защита проекта",
        # У шаблонов разная аттестация по умолчанию — фиксируем, чтобы
        # сравнивать только то, куда ушли часы лабораторных.
        attestation_hours=2,
    )


def test_general_program_has_no_lab_hours() -> None:
    program = assemble_program(_request(ProgramType.DOP_OBSH), WITH_LABS)

    assert program.modules_hours.labs == 0
    assert all(t.hours.labs == 0 for t in program.all_topics())
    assert program.grand_total == 36


def test_qualification_program_keeps_lab_hours() -> None:
    """В шаблоне повышения квалификации колонка ЛБ есть — её не трогаем."""
    program = assemble_program(_request(ProgramType.DPP_PK), WITH_LABS)

    assert program.modules_hours.labs > 0


def test_lab_hours_go_to_practice_not_self_study() -> None:
    """Лабораторная — занятие в классе: аудиторные часы не должны уменьшиться."""
    pk = assemble_program(_request(ProgramType.DPP_PK), WITH_LABS)
    obsh = assemble_program(_request(ProgramType.DOP_OBSH), WITH_LABS)

    assert obsh.modules_hours.contact == pk.modules_hours.contact
    assert obsh.modules_hours.self_study == pk.modules_hours.self_study


@pytest.mark.anyio
async def test_request_1_schedule_still_holds() -> None:
    """Чек-лист запроса 1: 8 занятий по 2 часа — все 16 часов в классе."""
    intake, _ = await run_intake(MockProvider(), REQUEST_1)
    request = intake.request
    apply_answers(request, {"audience": "школьники", "entry_level": "без опыта",
                            "attestation_form": "защита проекта"})
    program, _ = await build_skeleton(MockProvider(), request)
    hours = program.modules_hours

    assert hours.labs == 0
    assert hours.self_study == 0
    assert hours.contact + program.final_attestation.hours == 16
    assert hours.practice >= 8


def _plan_tables(data: bytes) -> list[list[str]]:
    """Шапки таблиц с часами — те, где есть строка «Итого (час.)»."""
    headers = []
    for table in Document(BytesIO(data)).tables:
        cells = " ".join(c.text for row in table.rows for c in row.cells)
        if "Итого (час.)" in cells:
            headers.append([c.text for c in table.rows[0].cells])
    return headers


def test_docx_general_program_matches_the_sample() -> None:
    program = assemble_program(_request(ProgramType.DOP_OBSH), WITH_LABS)

    headers = _plan_tables(render_docx(program))

    assert len(headers) == 2, "нужны учебный план и учебно-тематический план"
    for header in headers:
        assert "ЛБ" not in header
        assert "Всего аудиторных" in header
        assert "Самостоятельная работа" in header


def test_docx_qualification_program_matches_the_sample() -> None:
    program = assemble_program(_request(ProgramType.DPP_PK), WITH_LABS)

    headers = _plan_tables(render_docx(program))

    assert len(headers) == 2, "нужны учебный план и учебно-тематический план"
    for header in headers:
        assert ["ЛК", "ПР", "ЛБ", "СРС"] == [h for h in header if h in ("ЛК", "ПР", "ЛБ", "СРС")]


def test_general_program_totals_add_up_in_docx() -> None:
    """Аудиторные + самостоятельная = общая трудоёмкость в строке «Итого»."""
    program = assemble_program(_request(ProgramType.DOP_OBSH), WITH_LABS)
    document = Document(BytesIO(render_docx(program)))

    for table in document.tables:
        last = [c.text for c in table.rows[-1].cells]
        if "Итого (час.)" not in last:
            continue
        total, classroom, _lk, _pr, self_study = (float(x.replace(",", ".")) for x in last[2:7])
        assert classroom + self_study == total == 36
